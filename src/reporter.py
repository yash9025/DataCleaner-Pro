import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from pathlib import Path
import pandas as pd

class ExcelReportGenerator:
    def __init__(self, cleaned_df, audit_log, output_path="reports/Cleaned_Executive_Report.xlsx"):
        self.df = cleaned_df
        self.audit = audit_log
        self.output_path = Path(output_path)
        self.output_path.parent.mkdir(parents=True, exist_ok=True)
        self.wb = openpyxl.Workbook()
        # Remove default sheet
        self.wb.remove(self.wb.active)

    def _style_header_cell(self, cell, text, bg_color="1F4E79"):
        cell.value = text
        cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color=bg_color, end_color=bg_color, fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    def _build_kpi_card(self, ws, start_col, start_row, title, value, subtext="", card_color="2F5597"):
        """Creates a modern executive KPI card."""
        c1 = ws.cell(row=start_row, column=start_col)
        c2 = ws.cell(row=start_row+1, column=start_col)
        c3 = ws.cell(row=start_row+2, column=start_col)
        
        c1.value = title.upper()
        c1.font = Font(name="Calibri", size=9, bold=True, color="595959")
        c1.alignment = Alignment(horizontal="center", vertical="center")

        c2.value = value
        c2.font = Font(name="Calibri", size=18, bold=True, color=card_color)
        c2.alignment = Alignment(horizontal="center", vertical="center")

        c3.value = subtext
        c3.font = Font(name="Calibri", size=8, italic=True, color="7F7F7F")
        c3.alignment = Alignment(horizontal="center", vertical="center")

        # Border styling for card
        thin = Side(border_style="thin", color="D9D9D9")
        card_fill = PatternFill(start_color="F2F4F7", end_color="F2F4F7", fill_type="solid")
        for r in range(start_row, start_row+3):
            cell = ws.cell(row=r, column=start_col)
            cell.fill = card_fill
            cell.border = Border(top=thin, left=thin, right=thin, bottom=thin)

    def generate_summary_sheet(self):
        ws = self.wb.create_sheet(title="Executive Summary")
        ws.views.sheetView[0].showGridLines = True

        # Main Title Banner
        ws.merge_cells("B2:G2")
        title_cell = ws["B2"]
        title_cell.value = "EXECUTIVE SALES & PERFORMANCE SUMMARY"
        title_cell.font = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
        title_cell.fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
        title_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[2].height = 36

        # Subtitle
        ws.merge_cells("B3:G3")
        sub_cell = ws["B3"]
        sub_cell.value = f"Automated Pipeline Execution • Cleaned from {self.audit['file_name']} • Total Records: {len(self.df):,}"
        sub_cell.font = Font(name="Calibri", size=10, italic=True, color="595959")
        sub_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[3].height = 20

        # Calculations
        has_sales = 'Sales' in self.df.columns
        has_profit = 'Profit' in self.df.columns

        total_sales = float(self.df['Sales'].sum()) if has_sales else 0.0
        total_profit = float(self.df['Profit'].sum()) if has_profit else 0.0
        margin_pct = (total_profit / total_sales * 100) if total_sales > 0 else 0.0
        total_orders = len(self.df['Order ID'].unique()) if 'Order ID' in self.df.columns else len(self.df)

        # 4 KPI Cards in Row 5 to 7
        self._build_kpi_card(ws, start_col=2, start_row=5, title="Total Revenue", value=f"${total_sales:,.2f}", subtext="Cleaned gross sales", card_color="1F4E79")
        self._build_kpi_card(ws, start_col=3, start_row=5, title="Total Profit", value=f"${total_profit:,.2f}", subtext="Net operational profit", card_color="2E7D32" if total_profit>=0 else "C62828")
        self._build_kpi_card(ws, start_col=4, start_row=5, title="Profit Margin", value=f"{margin_pct:.1f}%", subtext="Margin on revenue", card_color="1565C0")
        self._build_kpi_card(ws, start_col=5, start_row=5, title="Unique Orders", value=f"{total_orders:,}", subtext="Deduplicated orders", card_color="6A1B9A")

        # Category Breakdown Table (Row 10)
        table_start_row = 10
        ws.cell(row=table_start_row, column=2, value="Category Performance Breakdown").font = Font(name="Calibri", size=13, bold=True, color="1F4E79")
        
        headers = ["Category", "Total Sales ($)", "Total Profit ($)", "Profit Margin", "Orders Count"]
        for col_idx, h in enumerate(headers, start=2):
            self._style_header_cell(ws.cell(row=table_start_row+1, column=col_idx), h, bg_color="2F5597")

        if 'Category' in self.df.columns:
            cat_group = self.df.groupby('Category').agg(
                sales=('Sales', 'sum') if has_sales else ('Category', 'count'),
                profit=('Profit', 'sum') if has_profit else ('Category', 'count'),
                orders=('Order ID', 'nunique') if 'Order ID' in self.df.columns else ('Category', 'count')
            ).reset_index()

            curr_row = table_start_row + 2
            for _, row_data in cat_group.iterrows():
                s_val = float(row_data['sales']) if has_sales else 0.0
                p_val = float(row_data['profit']) if has_profit else 0.0
                m_val = (p_val / s_val) if s_val > 0 else 0.0

                c_cat = ws.cell(row=curr_row, column=2, value=str(row_data['Category']))
                c_sales = ws.cell(row=curr_row, column=3, value=s_val)
                c_profit = ws.cell(row=curr_row, column=4, value=p_val)
                c_margin = ws.cell(row=curr_row, column=5, value=m_val)
                c_orders = ws.cell(row=curr_row, column=6, value=int(row_data['orders']))

                c_sales.number_format = "$#,##0.00"
                c_profit.number_format = "$#,##0.00"
                c_margin.number_format = "0.0%"
                c_orders.number_format = "#,##0"

                # Alternating light background
                row_fill = PatternFill(start_color="F9FAFB" if curr_row % 2 == 0 else "FFFFFF", end_color="F9FAFB" if curr_row % 2 == 0 else "FFFFFF", fill_type="solid")
                for c in [c_cat, c_sales, c_profit, c_margin, c_orders]:
                    c.fill = row_fill
                    c.border = Border(bottom=Side(style="thin", color="E0E0E0"))

                curr_row += 1

            # Summary Totals Row
            ws.cell(row=curr_row, column=2, value="TOTAL").font = Font(bold=True)
            t_sales = ws.cell(row=curr_row, column=3, value=f"=SUM(C{table_start_row+2}:C{curr_row-1})")
            t_profit = ws.cell(row=curr_row, column=4, value=f"=SUM(D{table_start_row+2}:D{curr_row-1})")
            t_margin = ws.cell(row=curr_row, column=5, value=f"=D{curr_row}/C{curr_row}")
            t_orders = ws.cell(row=curr_row, column=6, value=f"=SUM(F{table_start_row+2}:F{curr_row-1})")

            t_sales.number_format = "$#,##0.00"
            t_profit.number_format = "$#,##0.00"
            t_margin.number_format = "0.0%"
            t_orders.number_format = "#,##0"
            for c in [ws.cell(row=curr_row, column=2), t_sales, t_profit, t_margin, t_orders]:
                c.font = Font(bold=True)
                c.fill = PatternFill(start_color="E9EEF4", end_color="E9EEF4", fill_type="solid")
                c.border = Border(top=Side(style="thin", color="000000"), bottom=Side(style="double", color="000000"))

        # Set column widths nicely for Summary
        col_widths = {2: 24, 3: 20, 4: 20, 5: 18, 6: 18}
        for col_idx, width in col_widths.items():
            ws.column_dimensions[get_column_letter(col_idx)].width = width

    def generate_cleaned_data_sheet(self):
        ws = self.wb.create_sheet(title="Cleaned Data")
        ws.views.sheetView[0].showGridLines = True

        # Write Headers
        for col_idx, col_name in enumerate(self.df.columns, start=1):
            cell = ws.cell(row=1, column=col_idx)
            self._style_header_cell(cell, col_name, bg_color="1F4E79")

        # Write Rows
        thin_border = Border(
            left=Side(style='thin', color='E5E7EB'),
            right=Side(style='thin', color='E5E7EB'),
            top=Side(style='thin', color='E5E7EB'),
            bottom=Side(style='thin', color='E5E7EB')
        )

        for row_idx, row in enumerate(self.df.itertuples(index=False), start=2):
            is_even = (row_idx % 2 == 0)
            row_fill = PatternFill(start_color="F8FAFC" if is_even else "FFFFFF", fill_type="solid")

            for col_idx, val in enumerate(row, start=1):
                cell = ws.cell(row=row_idx, column=col_idx, value=val)
                cell.fill = row_fill
                cell.border = thin_border
                cell.font = Font(name="Calibri", size=10)

                col_name = self.df.columns[col_idx - 1]
                if col_name.lower() in ['sales', 'profit']:
                    cell.number_format = "$#,##0.00"
                    cell.alignment = Alignment(horizontal="right")
                elif col_name.lower() in ['discount']:
                    cell.number_format = "0.0%"
                    cell.alignment = Alignment(horizontal="right")
                elif col_name.lower() in ['quantity', 'postal code']:
                    cell.number_format = "#,##0"
                    cell.alignment = Alignment(horizontal="right")
                elif 'date' in col_name.lower():
                    cell.alignment = Alignment(horizontal="center")

        # Auto-filter and freeze header row
        ws.auto_filter.ref = f"A1:{get_column_letter(len(self.df.columns))}{len(self.df)+1}"
        ws.freeze_panes = "A2"

        # Auto-adjust column widths
        for col in ws.columns:
            col_letter = get_column_letter(col[0].column)
            max_len = max(len(str(cell.value or '')) for cell in col[:100])
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    def generate_audit_log_sheet(self):
        ws = self.wb.create_sheet(title="Cleaning Audit Log")
        ws.views.sheetView[0].showGridLines = True

        ws.cell(row=2, column=2, value="DATA INTEGRITY & AUTOMATION AUDIT REPORT").font = Font(name="Calibri", size=14, bold=True, color="1F4E79")
        ws.cell(row=3, column=2, value=f"Generated on {pd.Timestamp.now().strftime('%Y-%m-%d %H:%M:%S')}").font = Font(italic=True, color="7F7F7F")

        metrics = [
            ("Original File Name", self.audit["file_name"]),
            ("Original File Size", f"{self.audit['file_size_kb']} KB"),
            ("Raw Records Ingested", f"{self.audit['raw_rows']:,}"),
            ("Duplicate Records Removed", f"{self.audit['duplicates_removed']:,}"),
            ("Missing Values Treated", f"{self.audit['missing_values_before']:,}"),
            ("Final Cleaned Records", f"{self.audit['final_rows']:,}"),
            ("Data Retention Rate", f"{(self.audit['final_rows'] / self.audit['raw_rows'] * 100):.1f}%"),
            ("Total Execution Time", f"{self.audit['processing_time_sec']} seconds")
        ]

        # Table Header
        self._style_header_cell(ws.cell(row=5, column=2), "Metric / Step", bg_color="2F5597")
        self._style_header_cell(ws.cell(row=5, column=3), "Audit Result", bg_color="2F5597")

        curr_row = 6
        for label, val in metrics:
            c1 = ws.cell(row=curr_row, column=2, value=label)
            c2 = ws.cell(row=curr_row, column=3, value=val)
            c1.font = Font(bold=True)
            c2.alignment = Alignment(horizontal="right")
            fill = PatternFill(start_color="F9FAFB" if curr_row % 2 == 0 else "FFFFFF", fill_type="solid")
            c1.fill = fill
            c2.fill = fill
            c1.border = Border(bottom=Side(style="thin", color="E0E0E0"))
            c2.border = Border(bottom=Side(style="thin", color="E0E0E0"))
            curr_row += 1

        # Transformations Performed
        curr_row += 2
        ws.cell(row=curr_row, column=2, value="Column Transformations & Standardizations").font = Font(size=12, bold=True, color="1F4E79")
        curr_row += 1

        self._style_header_cell(ws.cell(row=curr_row, column=2), "#", bg_color="418AB3")
        self._style_header_cell(ws.cell(row=curr_row, column=3), "Transformation Applied", bg_color="418AB3")
        curr_row += 1

        for idx, step in enumerate(self.audit.get("cleaned_columns", []), start=1):
            c1 = ws.cell(row=curr_row, column=2, value=idx)
            c2 = ws.cell(row=curr_row, column=3, value=step)
            c1.alignment = Alignment(horizontal="center")
            curr_row += 1

        ws.column_dimensions["B"].width = 30
        ws.column_dimensions["C"].width = 50

    def save(self):
        self.generate_summary_sheet()
        self.generate_cleaned_data_sheet()
        self.generate_audit_log_sheet()
        self.wb.save(self.output_path)
        print(f"Executive Report saved to: {self.output_path.resolve()}")
        return self.output_path

if __name__ == "__main__":
    from cleaner import DataCleaner
    cleaner = DataCleaner("data/messy_sales_data.xlsx")
    cleaned_df, audit = cleaner.clean()
    reporter = ExcelReportGenerator(cleaned_df, audit, "reports/Cleaned_Executive_Report.xlsx")
    reporter.save()
